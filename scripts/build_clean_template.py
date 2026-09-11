"""Create a reusable, data-free template from the supplied manual workbook."""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

# Allow the documented `python scripts/build_clean_template.py ...` invocation.
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.services.excel_writer import SHEET_NAME, _header_row, find_row_containing, preserve_sheet_drawings, remove_external_links, replace_workbook_logo


def _clear_cell(cell) -> None:
    if not isinstance(cell, MergedCell):
        cell.value = None


def _is_variable_value(value: object) -> bool:
    if not isinstance(value, str):
        return False
    normalized = " ".join(value.lower().split())
    return (
        normalized.startswith("=")
        or bool(re.search(r"\bpo\s*0?27956\b|\b027956\b|\b019347\b|\b019458\b", normalized))
        or "code substitution" in normalized
        or "quantities update" in normalized
        or "included the item" in normalized
        or "net 240" in normalized
        or bool(re.search(r"\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)/20\d{2}\b", normalized))
        or bool(re.search(r"\b\d{2}/\d{2}/20\d{2}\b", normalized))
    )


def build_clean_template(reference: Path, output: Path, logo: Path | None = None) -> None:
    wb = load_workbook(reference, keep_links=False)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Aba {SHEET_NAME!r} ausente no workbook de referência")
    ws = wb[SHEET_NAME]
    item_start = _header_row(ws) + 1
    approval = find_row_containing(ws, "approval")
    if approval is None:
        raise ValueError("Bloco APPROVAL não encontrado no workbook de referência")
    # Items and their manual lookup/formula cells are always variable.
    for row in range(item_start, approval - 1):
        for cell in ws[row][:19]:
            _clear_cell(cell)
    # Clear only known variable values below the table. Fixed labels and formatting remain.
    for row in ws.iter_rows(min_row=approval):
        for cell in row:
            if _is_variable_value(cell.value):
                _clear_cell(cell)
    # Historical substitution/MOQ tables are visual structure only. Clear every
    # data row beneath their headers; never derive or copy those records.
    table_headers: list[int] = []
    for row in ws.iter_rows(min_row=approval):
        labels = " ".join(str(cell.value or "").lower() for cell in row)
        if ("old code" in labels and "new code" in labels) or ("material code" in labels and "moq" in labels):
            table_headers.append(row[0].row)
    for index, header in enumerate(table_headers):
        end = table_headers[index + 1] if index + 1 < len(table_headers) else ws.max_row + 1
        for row_number in range(header + 1, end):
            for cell in ws[row_number]:
                _clear_cell(cell)
    # Operational locations in the manual workbook are not confirmed PDF fields.
    for row in ws.iter_rows(min_row=approval):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.lower().startswith(("place of loading:", "place of delivery:")):
                _clear_cell(cell)
    # Formulae elsewhere in the material sheet are never part of the MVP output.
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                _clear_cell(cell)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()
    remove_external_links(output)
    preserve_sheet_drawings(reference, output)
    if logo is not None:
        replace_workbook_logo(output, logo)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reference", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--logo", type=Path, default=Path("assets/logo eletra.png"))
    args = parser.parse_args()
    build_clean_template(args.reference, args.output, args.logo)


if __name__ == "__main__":
    main()
