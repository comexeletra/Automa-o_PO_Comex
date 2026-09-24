from datetime import date
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from app.models import PurchaseOrder, PurchaseOrderItem
from app.services.excel_writer import _translate_merges, write_purchase_order


def _template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Material for production"
    headers = ["PO #", "#", "Supply code", "NCM", "Eletra code", "Product description", "", "", "", "", "", "", "Total QTY", "Unit price", "Total (CNY)", "Finalidade", "Centro de custos", "Requer LI?", "Projeto"]
    for column, value in enumerate(headers, 1): ws.cell(18, column).value = value
    ws.merge_cells("F18:L18"); ws["F18"] = "Product description"
    ws.merge_cells("A19:A20"); ws.merge_cells("A21:L21"); ws.merge_cells("N21:O21")
    ws["A21"] = "TOTAL"; ws["A23"] = "APPROVAL"; ws["A25"] = "REMARKS"; ws["A27"] = "PO"; ws["A28"] = "SC"; ws["A29"] = "Payment Terms"
    wb.save(path)


def _po(count: int = 3) -> PurchaseOrder:
    items = tuple(PurchaseOrderItem("001", n, f"2010100{n:03d}", f"Description {n}", "UN", Decimal("2.000"), Decimal("1.2345000"), Decimal("2.47"), date(2026, 6, 16), None, "019347") for n in range(1, count + 1))
    return PurchaseOrder("000123", "Buyer", "12.115.480/0001-15", None, "Seller", "EX000135", None, date(2026, 6, 16), "240 DIAS", Decimal("7.41"), items)


def test_writer_resizes_and_keeps_lower_sections(tmp_path):
    template, output = tmp_path / "template.xlsx", tmp_path / "result.xlsx"
    _template(template)
    write_purchase_order(_po(), template, output)
    ws = load_workbook(output, data_only=False)["Material for production"]
    assert [ws.cell(row, 3).value for row in range(19, 22)] == ["2010100001", "2010100002", "2010100003"]
    assert ws["A24"].value == "APPROVAL"
    assert ws["A26"].value == "REMARKS"
    assert ws["D19"].value is None and ws["P19"].value is None
    with ZipFile(output) as archive:
        assert not any(name.startswith("xl/externalLinks/") for name in archive.namelist())


def test_writer_shrinks_item_area(tmp_path):
    template, output = tmp_path / "template.xlsx", tmp_path / "result.xlsx"
    _template(template)
    write_purchase_order(_po(1), template, output)
    ws = load_workbook(output)["Material for production"]
    assert ws["A22"].value == "APPROVAL"


def test_compact_layout_hides_internal_only_fields(tmp_path):
    template, output = tmp_path / "template.xlsx", tmp_path / "result.xlsx"
    _template(template)
    write_purchase_order(_po(1), template, output, compact_layout=True)
    ws = load_workbook(output)["Material for production"]

    assert all(ws.column_dimensions[column].hidden for column in ("E", "P", "Q", "R", "S"))
    assert ws["E18"].value is None
    assert all(ws[f"{column}18"].value is None for column in ("P", "Q", "R", "S"))


def test_translate_merges_recovers_missing_stale_merge_cell():
    """A stale merge after row movement must not raise KeyError in openpyxl."""
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A3:C3")
    del ws._cells[(3, 2)]  # Same incomplete internal state as the production failure.

    _translate_merges(ws, item_start=10, lower_start=20, total_row=19, new_count=1, delta=0)

    assert "A3:C3" in {str(merged) for merged in ws.merged_cells.ranges}


def test_writer_extends_item_area_before_total_without_currency_format(tmp_path):
    output = tmp_path / "result.xlsx"
    write_purchase_order(_po(206), Path("app/resources/po_template.xlsx"), output)
    ws = load_workbook(output, data_only=False)["Material for production"]
    # The original template supports 193 rows. Extra rows must use the item-row
    # style rather than the total/approval rows' currency format.
    assert ws["B212"].value == 194
    assert ws["B212"].number_format == "0"
    assert ws["B212"].font.sz == ws["B19"].font.sz == 14
    assert ws["B212"]._style.borderId == ws["B19"]._style.borderId
    assert ws["C214"].font.sz == ws["C19"].font.sz == 14
    assert ws["C214"]._style.borderId == ws["C19"]._style.borderId
    assert ws["B224"].value == 206
    assert Decimal(str(ws["N225"].value)) == Decimal("7.41")
    assert ws["A227"].value == "APPROVAL"
    # The signature block's custom row heights must move together with its
    # merged cells and drawings; otherwise the section becomes visually skewed.
    assert ws.row_dimensions[229].height is None
    assert ws.row_dimensions[242].height == 15.75
    for cell in ("I231", "J231", "K231", "I238", "J238"):
        assert ws[cell].border.top.style == "thin"
