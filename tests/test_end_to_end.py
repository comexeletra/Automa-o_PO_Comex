from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.services.processor import process_purchase_order


PDF = Path("samples/sample_po_027956.pdf")
TEMPLATE = Path("templates/po_template.xlsx")


@pytest.mark.skipif(not PDF.exists() or not TEMPLATE.exists(), reason="Arquivos de referência não estão disponíveis")
def test_reference_pdf_to_excel_end_to_end(tmp_path):
    output = tmp_path / "PO_027956.xlsx"
    result = process_purchase_order(PDF, output, TEMPLATE)
    assert len(result.po.items) == 193
    wb = load_workbook(output, data_only=False, keep_links=False)
    ws = wb["Material for production"]
    assert ws["A19"].value == "027956"
    assert ws["B19"].value == 1
    assert ws["C19"].value == "2010100040"
    assert ws["D19"].value is None and ws["E19"].value is None
    assert ws["F19"].value.startswith("RESISTOR 390K")
    assert ws["M19"].value == 75000
    assert ws["N19"].value == 0.010416
    assert ws["O19"].value == 781.2
    assert all(ws.cell(row, col).value is None for row in range(19, 212) for col in (4, 5, 16, 17, 18, 19))
    assert ws["M212"].value == 53389684
    assert ws["N212"].value == 9361869.5
    assert any("APPROVAL" in str(cell.value) for cell in ws[214])
    assert any("REMARKS" in str(cell.value) for cell in ws[229])
    assert ws["A233"].value == "PO027956"
    assert ws["A234"].value == "SC - 019347 - 019458"
    assert ws["A236"].value == "Payment Terms: 240 DIAS"
    # Lower tables remain as visual structure, with no inherited manual records.
    assert all(ws.cell(row, column).value is None for row in range(240, 245) for column in range(1, 7))
    assert all(ws.cell(row, column).value is None for row in range(247, 259) for column in range(1, 7))
    wb.close()
    with ZipFile(output) as archive:
        assert not any(name.startswith("xl/externalLinks/") for name in archive.namelist())
        assert "xl/drawings/drawing1.xml" in archive.namelist()
        assert "xl/media/image1.png" in archive.namelist()
